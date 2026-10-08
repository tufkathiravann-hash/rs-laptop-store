import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=50, bottom=50, left=50, right=50):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_double_border_to_section(section):
    # Set page size to US Letter
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    sectPr = section._sectPr
    # Add page borders: double line border
    pgBorders = parse_xml(
        '<w:pgBorders %s w:offsetFrom="page">\n'
        '  <w:top w:val="double" w:sz="12" w:space="24" w:color="000000"/>\n'
        '  <w:left w:val="double" w:sz="12" w:space="24" w:color="000000"/>\n'
        '  <w:bottom w:val="double" w:sz="12" w:space="24" w:color="000000"/>\n'
        '  <w:right w:val="double" w:sz="12" w:space="24" w:color="000000"/>\n'
        '</w:pgBorders>' % nsdecls('w')
    )
    sectPr.append(pgBorders)

def add_footer_page_number(doc, start_page=3):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # In Word, page number field
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    
    # We can set starting page number in section to 3
    sectPr = section._sectPr
    pgNumType = OxmlElement('w:pgNumType')
    pgNumType.set(qn('w:start'), str(start_page))
    sectPr.append(pgNumType)

    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p._element.append(fldSimple)

def build_word_report(filename):
    doc = Document()
    add_double_border_to_section(doc.sections[0])
    add_footer_page_number(doc, start_page=3)

    sc_dir = r'd:\RS PROJECT\PROJECT REPORT\screenshots\processed'

    # Styles helper
    def add_para(text, font_size=10, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=4, indent=0.25):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if indent > 0:
            p.paragraph_format.first_line_indent = Inches(indent)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_heading(text, space_before=8, space_after=3):
        return add_para(text, font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=space_before, space_after=space_after, indent=0)

    def add_bullet(text, space_after=2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run("•  " + text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_caption(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(9.5)
        run.font.bold = True
        return p

    def add_img(img_path, width_in=5.8):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(width_in))

    # ================= PAGE 1 =================
    add_para("CHAPTER - 2", font_size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=2, indent=0)
    add_para("DAILY PROGRESS", font_size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=12, indent=0)

    add_heading("2.1 Project Idea and Requirement Analysis", space_before=4)
    add_para(
        "The primary objective of RS E-Commerce is to create a specialized, high-performance online laptop shopping platform that eliminates the complexity and friction of purchasing premium computing devices. Unlike generic retail platforms, purchasing modern laptops demands in-depth technical specifications, transparent configuration options, brand categorization, and a streamlined purchasing workflow. Requirement analysis identified key system modules including dynamic catalog browsing, multi-criteria hardware filtering (CPU, RAM, GPU, SSD), real-time cart state management, simulated order checkout, and comprehensive customer enquiry mechanisms tailored specifically for modern laptop consumers."
    )

    add_heading("2.2 Website Structure", space_before=6)
    add_para(
        "The RS E-Commerce web application is designed using a modular, scalable component-driven architecture built with React 18, TypeScript, and Vite. The system maintains strict separation of concerns through dedicated modular directories: /src/pages for primary route views, /src/components for reusable UI controls, /src/context for global state providers (CartContext, AuthContext, WishlistContext), /src/data for structured catalog datasets, and /src/types for strict TypeScript contracts. Tailwind CSS delivers a uniform dark theme design system with responsive layouts across all screen resolutions."
    )

    add_img(os.path.join(sc_dir, 'fig_2_1_structure.png'), width_in=5.7)
    add_caption("Fig:2.1 Website Structure")
    doc.add_page_break()

    # ================= PAGE 2 =================
    add_heading("2.3 Laptop Product Section", space_before=0)
    add_para(
        "The laptop product section displays available computers in a dynamic, responsive grid layout. Products are rendered dynamically through React components mapping over typed catalog models rather than being static HTML markup. Each product card presents:",
        indent=0
    )
    add_bullet("Laptop product image with high-resolution thumbnail")
    add_bullet("Laptop model name and manufacturer brand")
    add_bullet("Product category badge and hardware specifications")
    add_bullet("Current price, original price, and active discounts")
    add_bullet("Add to Cart button and quick wishlist toggle")

    add_img(os.path.join(sc_dir, 'fig_2_2_product_section.png'), width_in=5.5)
    add_caption("Fig:2.2 Laptop Product Section")

    add_heading("2.4 Developing the Home Page", space_before=6)
    add_para(
        "The Home Page serves as the central showcase and primary interactive gateway for the RS E-Commerce laptop platform. It incorporates a modern cyberpunk-inspired hero banner highlighting flagship machines, an intuitive quick-search input, manufacturer brand carousels, top promotional deals, and verified buyer reviews. Robust navigation headers enable instant routing between catalog categories, shopping cart summaries, and user account portals."
    )
    doc.add_page_break()

    # ================= PAGE 3 =================
    add_img(os.path.join(sc_dir, 'fig_2_3_homepage.png'), width_in=5.8)
    add_caption("Fig:2.3 Developing the Home Page")

    add_heading("2.5 Laptop Brand Categories", space_before=6)
    add_para(
        "Users can filter laptops seamlessly using dedicated brand category buttons. The implemented brand categories include:",
        indent=0
    )
    add_bullet("ASUS (ROG Strix, TUF Gaming, and ZenBook series)")
    add_bullet("Apple (MacBook Pro and MacBook Air M-Series)")
    add_bullet("Dell (Alienware and XPS performance series)")
    add_bullet("Lenovo (Legion gaming and ThinkPad series)")
    add_bullet("HP (Omen gaming and Spectre series)")

    add_img(os.path.join(sc_dir, 'fig_2_4_brand_categories.png'), width_in=5.8)
    add_caption("Fig:2.4 Laptop Brand Categories")

    add_heading("2.6 Search Function", space_before=6)
    add_para(
        "The real-time search functionality enables customers to instantly discover specific laptop models, processors, or hardware keywords. Whenever a user types into the search bar, React state triggers instantaneous filtering across titles, descriptions, and technical specifications, displaying matching laptop results without page reloads."
    )
    doc.add_page_break()

    # ================= PAGE 4 =================
    add_heading("2.7 Laptop Search and Filtering", space_before=0)
    add_para(
        "The multi-faceted search and filtering system allows customers to refine laptop selections according to specific technical criteria. Users can filter products by manufacturer brand, price range slider, RAM capacity (16GB, 32GB, 64GB), storage size (512GB, 1TB, 2TB SSD), and primary usage category (Gaming, Workstation, Ultralight, Creator). Live counters display matching models, enabling quick comparison."
    )

    add_heading("2.8 Shopping Cart", space_before=6)
    add_para(
        "The shopping cart represents a crucial interactive module of the website. Built using React Context (CartContext) with browser localStorage persistence, it tracks selected laptops, custom warranty add-ons, and quantities. When a customer adds an existing product, the system automatically increments the item quantity rather than creating duplicate entries, providing a smooth user experience."
    )

    add_img(os.path.join(sc_dir, 'fig_2_5_shopping_cart.png'), width_in=5.7)
    add_caption("Fig:2.5 Shopping Cart")
    doc.add_page_break()

    # ================= PAGE 5 =================
    add_heading("2.9 Laptop Product Listing", space_before=0)
    add_para(
        "The product listing page presents the comprehensive catalog of laptops with rich visuals, detailed specification badges, live stock availability, customer review ratings, and direct Add to Cart triggers. Shoppers can sort products by price, popularity, release date, and user ratings to discover laptops tailored to their performance and budget requirements."
    )

    add_heading("2.10 Laptop Product Details", space_before=10)
    add_para(
        "The product details page delivers an exhaustive technical overview of each laptop model. It incorporates high-resolution multi-angle image galleries, detailed hardware specification tables (display refresh rates, GPU TGP wattage, thermal cooling architecture, port selection), warranty packages, customer reviews, and immediate Add to Cart or Buy Now purchasing options."
    )

    add_heading("2.11 Wishlist", space_before=10)
    add_para(
        "The Wishlist allows customers to bookmark desired laptops for later review and purchase. Shoppers can save premium configurations without immediately adding them to their active shopping cart. Customers can revisit their saved wishlist anytime and effortlessly transfer selected laptops into the cart when ready to purchase."
    )

    add_heading("2.12 User Registration and Login", space_before=10)
    add_para(
        "RS E-Commerce provides a streamlined authentication module for customer accounts. Users can register and sign in using basic profile credentials including full name, email address, contact phone number, and password. Authenticated customers gain access to personalized order history, saved shipping addresses, and persistent cart synchronisation across browsing sessions."
    )
    doc.add_page_break()

    # ================= PAGE 6 =================
    add_heading("2.13 Quantity Management", space_before=0)
    add_para("Users can easily adjust the quantity of laptops directly within the shopping cart drawer and cart page using intuitive + and - buttons.", indent=0)
    add_para("For example:", indent=0, space_after=2)
    add_bullet("Product quantity: 1")
    add_para("Clicking + changes it to:", indent=0, space_after=2)
    add_bullet("Product quantity: 2")
    add_para("Clicking - decreases it again.", indent=0, space_after=6)

    add_heading("2.14 Payment Options", space_before=10)
    add_para(
        "RS E-Commerce is structured to support flexible and secure payment methods. Supported checkout options include Credit Cards, Debit Cards, UPI (Google Pay, PhonePe, Paytm), Net Banking across major banks, No-Cost EMI financing options for high-value laptops, and Cash on Delivery (COD) for supported delivery zones. The checkout simulation verifies all transaction fields cleanly."
    )

    add_heading("2.15 Order Confirmation", space_before=10)
    add_para(
        "Upon successfully completing the checkout flow, customers receive an immediate order confirmation summary. The confirmation screen presents a unique order ID, purchased laptop models, billing breakdown, shipping address, chosen payment method, and estimated delivery timeline, with options to download an invoice or track order progress."
    )

    add_heading("2.16 Order Tracking", space_before=10)
    add_para(
        "The order tracking module allows buyers to monitor shipment milestones from processing to doorstep delivery. Tracking stages include Order Confirmed, Quality Inspection, Dispatched, In Transit, Out for Delivery, and Delivered. A transparent tracking timeline establishes buyer confidence and reduces customer support inquiries."
    )
    doc.add_page_break()

    # ================= PAGE 7 =================
    add_heading("2.17 Delivery Service", space_before=0)
    add_para(
        "Reliable logistics and secure transit are essential when purchasing premium laptops. RS E-Commerce integrates automated pincode verification, insured courier handling, and scheduled delivery timeframes. Laptops are dispatched in shock-resistant packaging with delivery tracking numbers to ensure safe and timely arrival."
    )

    add_heading("2.18 Cart Total Calculation", space_before=10)
    add_para("The website automatically calculates the total price dynamically.", indent=0)
    add_para("The calculation used is:", indent=0, space_after=2)
    add_bullet("Total = Product Price × Quantity")
    add_para("For multiple products:", indent=0, space_after=2)
    add_bullet("Cart Total = Σ (Product Price × Quantity)")
    add_para("The total is dynamically updated whenever:", indent=0, space_after=2)
    add_bullet("A product is added")
    add_bullet("Product quantity is changed")
    add_bullet("A product is removed")
    add_para("Prices are formatted using the Indian numbering system through:", indent=0, space_after=2)
    add_bullet('toLocaleString("en-IN")')

    add_heading("2.19 Checkout Function", space_before=10)
    add_para(
        "The website includes a comprehensive checkout button. When the user clicks Proceed to Checkout, JavaScript validates whether the cart contains items. If products are available, the website navigates to the checkout interface for shipping address entry, order review, and payment option selection."
    )
    doc.add_page_break()

    # ================= PAGE 8 =================
    add_img(os.path.join(sc_dir, 'fig_2_6_checkout.png'), width_in=5.8)
    add_caption("Fig:2.6 Checkout Function")

    add_heading("2.20 Contact and Enquiry Section", space_before=6)
    add_para("The website features a dedicated Contact Us and technical enquiry section. The form includes:", indent=0)
    add_bullet("Name")
    add_bullet("Email")
    add_bullet("Phone number")
    add_bullet("Message / Technical Enquiry")
    add_bullet("Contact button")

    add_img(os.path.join(sc_dir, 'fig_2_7_contact.png'), width_in=5.8)
    add_caption("Fig:2.7 Contact and Enquiry Section")
    doc.add_page_break()

    # ================= PAGE 9 =================
    add_heading("2.21 Customer Reviews and Ratings", space_before=0)
    add_para(
        "Customer reviews and star ratings help prospective buyers make confident decisions. Buyers can submit verified feedback regarding build quality, battery endurance, thermal performance, and overall satisfaction. RS E-Commerce displays average ratings and verified reviews on product pages to foster community trust."
    )

    add_heading("2.22 Special Offers and Discounts", space_before=10)
    add_para(
        "Special offers and promotional campaigns are effective methods for boosting user engagement and conversions. RS E-Commerce provides discounts during seasonal events, festivals, tech expos, and product launches:",
        indent=0
    )
    add_bullet("Flat percentage discounts on select laptop brands")
    add_bullet("Bundle deals with gaming accessories and warranties")
    add_bullet("Free express delivery on prepaid orders")
    add_bullet("First-order student discount coupons")
    add_bullet("Brand-specific clearance discounts")

    add_heading("2.23 Promotional Laptop Deals", space_before=10)
    add_para(
        "The Deals section showcases limited-time offers featuring substantial price markdowns on top-tier gaming and workstation laptops. Real-time discount tags and promotional countdown timers create purchasing urgency and encourage immediate checkout."
    )

    add_heading("2.24 New Laptop Arrivals", space_before=10)
    add_para(
        "The New Arrivals section highlights newly released laptop models equipped with next-generation processors and dedicated GPUs. Customers can easily discover the latest hardware innovations with detailed launch specifications and direct purchase options."
    )
    doc.add_page_break()

    # ================= PAGE 10 =================
    add_heading("2.25 Best Selling Laptops", space_before=0)
    add_para(
        "The Best Sellers section highlights the most popular and highly rated laptop models based on sales volume, customer reviews, and benchmark reliability. Featuring verified customer favorites helps new shoppers identify dependable models across gaming, creator, and productivity categories."
    )

    add_heading("2.26 Responsive Web Design", space_before=8)
    add_para(
        "RS E-Commerce is developed with a mobile-first, responsive architecture built with Tailwind CSS. The user interface seamlessly adapts across mobile phones, tablets, laptops, and ultra-wide desktop monitors, featuring touch-friendly buttons, adaptive product grids, and smooth drawer navigation menus."
    )

    add_heading("2.27 Customer Support", space_before=8)
    add_para(
        "RS E-Commerce delivers comprehensive customer support covering technical specifications, order status, warranty coverage, and return queries. Support channels include hardware FAQs, email ticketing, and direct enquiry forms, ensuring prompt assistance and customer retention."
    )

    add_heading("2.28 Working Flow of the Website", space_before=8)
    add_para("The general workflow is:", indent=0, space_after=4)
    add_para("Open Website", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("Display Products", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("Search or Filter Products", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("Select Product", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    doc.add_page_break()

    # ================= PAGE 11 =================
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_before=0, space_after=2)
    add_para("Add to Cart", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("View Cart", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("Change Quantity / Remove Product", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("Calculate Total", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("↓", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=2)
    add_para("Display Checkout Message", align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, space_after=10)

    add_heading("2.29 Marketing and Customer Engagement", space_before=10)
    add_para(
        "RS E-Commerce utilizes digital marketing and customer engagement strategies to attract tech enthusiasts and maintain strong relationships with existing buyers. Marketing initiatives include social media campaigns, promotional banners, seasonal discount coupons, hardware comparison guides, and loyalty reward programs for repeat customers."
    )

    add_heading("2.30 Future Development", space_before=10)
    add_para(
        "RS E-Commerce can be expanded with advanced future enhancements such as AI-powered laptop recommendation wizards, 3D interactive model previewers, native mobile applications for iOS and Android, automated live chat assistance, genuine payment gateway integrations (Razorpay/Stripe), multi-currency support, and enterprise B2B bulk procurement portals."
    )

    doc.save(filename)
    print("Word document built successfully:", filename)

if __name__ == '__main__':
    out_docx = r'd:\RS PROJECT\PROJECT REPORT\RS_ECommerce_Daily_Progress_Report.docx'
    build_word_report(out_docx)
